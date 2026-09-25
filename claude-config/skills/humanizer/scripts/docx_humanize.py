"""Extract and rewrite the prose paragraphs of a .docx without touching its formatting.

    python docx_humanize.py extract makalah.docx paragraphs.json
    python docx_humanize.py apply   makalah.docx edits.json makalah-humanized.docx [--flatten]

`extract` lists every paragraph in the body (tables included) with its index, style,
section heading, and text. Italic and bold inside a paragraph are shown as *x* and **x**.
Each paragraph gets a `kind`:

    simple   plain text runs only; `apply` can rewrite it
    complex  has a citation field, footnote, image, hyperlink, equation, tracked change,
             or mixed fonts/sizes/underline; `apply` refuses it
    empty    no text

`edits.json` maps an index to the new text: {"12": "Teks baru dengan *istilah asing*."}.
`null` deletes that paragraph: {"14": null}. Indices always refer to the original file.
Keep the *x* / **x** marks where the formatting should survive. `apply` keeps the
paragraph's style, alignment, spacing and base font, rebuilds only the runs, and warns
when a number, year, (Author, 2020) citation, acronym, or name in the old text is missing
from the new. `--flatten` also rewrites paragraphs whose only problem is mixed fonts or
sizes (common after copy-paste), giving the whole paragraph the font of its main text.

Needs python-docx (`pip install python-docx`). Never overwrites the input file.
"""

import copy
import json
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

# Paragraph children that are safe to drop and rebuild around.
PARA_OK = {qn("w:pPr"), qn("w:r"), qn("w:proofErr"), qn("w:bookmarkStart"),
           qn("w:bookmarkEnd"), qn("w:permStart"), qn("w:permEnd")}
RUN_OK = {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:lastRenderedPageBreak")}
# Run properties that the *x* / **x** marks carry, or that never change how text looks.
MARKED = {qn("w:b"), qn("w:bCs"), qn("w:i"), qn("w:iCs")}
IGNORED = {qn("w:lang"), qn("w:noProof"), qn("w:rtl"), qn("w:webHidden"), qn("w:szCs")}

REASONS = {
    qn("w:hyperlink"): "hyperlink", qn("w:fldSimple"): "field", qn("w:sdt"): "content control",
    qn("w:ins"): "tracked change", qn("w:del"): "tracked change", qn("m:oMath"): "equation",
    qn("m:oMathPara"): "equation", qn("w:smartTag"): "smart tag",
    qn("w:fldChar"): "field (citation, TOC or page number)", qn("w:instrText"): "field",
    qn("w:footnoteReference"): "footnote", qn("w:endnoteReference"): "endnote",
    qn("w:drawing"): "image", qn("w:pict"): "image", qn("w:object"): "embedded object",
    qn("w:sym"): "symbol", qn("w:commentReference"): "comment",
}


def is_on(rpr, tag):
    el = rpr.find(tag) if rpr is not None else None
    return el is not None and el.get(qn("w:val")) not in ("0", "false", "none")


def base_key(rpr):
    """Run formatting with bold/italic removed, so runs can be compared."""
    if rpr is None:
        return ""
    r = copy.deepcopy(rpr)
    for child in list(r):
        if child.tag in MARKED or child.tag in IGNORED:
            r.remove(child)
    for attr in list(r.attrib):
        if "rsid" in attr:
            del r.attrib[attr]
    fonts = r.find(qn("w:rFonts"))
    if fonts is not None:
        fonts.attrib.pop(qn("w:hint"), None)
        if not fonts.attrib:
            r.remove(fonts)
    return etree.tostring(r) if len(r) or r.attrib else ""


def classify(p, flatten=False):
    for child in p:
        if child.tag not in PARA_OK:
            return "complex", REASONS.get(child.tag, child.tag.split("}")[-1])
    runs = p.findall(qn("w:r"))
    for r in runs:
        for child in r:
            if child.tag not in RUN_OK:
                return "complex", REASONS.get(child.tag, child.tag.split("}")[-1])
    keys = {base_key(r.find(qn("w:rPr"))) for r in runs if r.find(qn("w:t")) is not None}
    if len(keys) > 1 and not flatten:
        return "complex", "mixed fonts, sizes, colours or underline"
    text = run_text(p, marks=False)
    if not text.strip():
        return "empty", ""
    if "*" in text:
        return "complex", "literal asterisk in text"
    return "simple", ""


def run_text(p, marks=True):
    # Word splits text into many runs; merge neighbours with the same bold/italic first,
    # so "*istilah asing*" comes out as one marked span, not "*istilah* *asing*".
    segs = []
    for r in p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        s = ""
        for child in r:
            if child.tag == qn("w:t"):
                s += child.text or ""
            elif child.tag == qn("w:tab"):
                s += "\t"
            elif child.tag == qn("w:br"):
                s += "\n"
        if not s:
            continue
        flags = (is_on(rpr, qn("w:b")), is_on(rpr, qn("w:i"))) if marks else (False, False)
        if segs and segs[-1][1] == flags:
            segs[-1][0] += s
        else:
            segs.append([s, flags])
    out = []
    for s, (bold, italic) in segs:
        if (bold or italic) and s.strip():
            lead, core, trail = re.match(r"(\s*)(.*?)(\s*)$", s, re.S).groups()
            if bold:
                core = f"**{core}**"
            if italic:
                core = f"*{core}*"
            s = lead + core + trail
        out.append(s)
    return "".join(out)


def paragraphs(doc):
    body = doc.element.body
    return [p for p in body.iter(qn("w:p"))]


def style_of(p, doc):
    ppr = p.find(qn("w:pPr"))
    sid = ppr.find(qn("w:pStyle")).get(qn("w:val")) if ppr is not None and ppr.find(qn("w:pStyle")) is not None else None
    if sid is None:
        return "Normal"
    for s in doc.styles:
        if s.style_id == sid:
            return s.name
    return sid


def in_table(p):
    return any(a.tag == qn("w:tbl") for a in p.iterancestors())


def extract(src, dst):
    doc = Document(src)
    items, section = [], ""
    for i, p in enumerate(paragraphs(doc)):
        kind, reason = classify(p)
        style = style_of(p, doc)
        text = run_text(p) if kind == "simple" else run_text(p, marks=False)
        if style.lower().startswith(("heading", "judul", "title")) and text.strip():
            section = text.strip()
        item = {"i": i, "style": style, "section": section, "kind": kind, "text": text}
        if reason:
            item["reason"] = reason
        if in_table(p):
            item["in_table"] = True
        items.append(item)
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=1)
    counts = {k: sum(1 for x in items if x["kind"] == k) for k in ("simple", "complex", "empty")}
    print(f"{len(items)} paragraphs: {counts['simple']} simple, {counts['complex']} complex, {counts['empty']} empty -> {dst}")


def parse_marks(text):
    """Split 'a *b* **c**' into [(text, bold, italic), ...]."""
    segs, bold, italic, buf, i = [], False, False, "", 0
    while i < len(text):
        if text.startswith("**", i):
            segs.append((buf, bold, italic)); buf = ""; bold = not bold; i += 2
        elif text[i] == "*":
            segs.append((buf, bold, italic)); buf = ""; italic = not italic; i += 1
        else:
            buf += text[i]; i += 1
    segs.append((buf, bold, italic))
    if bold or italic:
        raise ValueError("unbalanced * or ** marks")
    return [s for s in segs if s[0]]


def facts(text):
    t = text.replace("*", "")
    nums = set(re.findall(r"\d[\d.,]*\d|\d", t))
    cites = set(re.findall(r"\([^()]*?\d{4}[a-z]?\)", t))
    acronyms = set(re.findall(r"\b[A-Z]{2,}\b", t))
    # Capitalised words mid-sentence are usually names: Kasmir, Jakarta, Undang-Undang.
    names = set(re.findall(r"(?<=[A-Za-z0-9,;] )[A-Z][a-z]+(?:[- ][A-Z][a-z]+)*", t))
    return nums | cites | acronyms | names


def make_run(base_rpr, text, bold, italic):
    r = etree.Element(qn("w:r"))
    rpr = copy.deepcopy(base_rpr) if base_rpr is not None else etree.Element(qn("w:rPr"))
    for child in list(rpr):
        if child.tag in MARKED:
            rpr.remove(child)
    # w:b and w:i sit near the top of rPr in schema order, after rStyle/rFonts.
    pos = sum(1 for c in rpr if c.tag in (qn("w:rStyle"), qn("w:rFonts")))
    for on, tags in ((italic, ("w:iCs", "w:i")), (bold, ("w:bCs", "w:b"))):
        if on:
            for tag in tags:
                rpr.insert(pos, etree.Element(qn(tag)))
    if len(rpr):
        r.append(rpr)
    parts = re.split(r"(\t|\n)", text)
    for part in parts:
        if part == "\t":
            r.append(etree.Element(qn("w:tab")))
        elif part == "\n":
            r.append(etree.Element(qn("w:br")))
        elif part:
            t = etree.SubElement(r, qn("w:t"))
            t.text = part
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return r


def apply(src, edits_path, dst, flatten=False):
    if os.path.normcase(os.path.abspath(dst)) == os.path.normcase(os.path.abspath(src)):
        sys.exit("refusing to overwrite the input file; give a new output name")
    doc = Document(src)
    ps = paragraphs(doc)
    with open(edits_path, encoding="utf-8") as f:
        edits = {int(k): v for k, v in json.load(f).items()}
    done, warnings, words_old, words_new = 0, [], 0, 0
    for i, new in sorted(edits.items()):
        p = ps[i]
        kind, reason = classify(p, flatten)
        if kind != "simple":
            warnings.append(f"[{i}] skipped: {kind} ({reason})")
            continue
        old = run_text(p)
        words_old += len(old.split())
        if new is None:
            # null deletes the paragraph. A table cell must keep at least one paragraph.
            cell = p.getparent()
            if cell.tag == qn("w:tc") and len(cell.findall(qn("w:p"))) == 1:
                warnings.append(f"[{i}] skipped: last paragraph in a table cell cannot be deleted")
                words_old -= len(old.split())
                continue
            lost = facts(old)
            if lost:
                warnings.append(f"[{i}] check: deleted paragraph held {', '.join(sorted(lost))}")
            cell.remove(p)
            done += 1
            continue
        words_new += len(new.replace("*", "").split())
        try:
            segs = parse_marks(new)
        except ValueError as e:
            warnings.append(f"[{i}] skipped: {e}")
            continue
        lost = facts(old) - facts(new)
        if lost:
            warnings.append(f"[{i}] check: {', '.join(sorted(lost))} is in the old text but not the new")
        # The run holding the most text sets the font for the rewrite. Only differs from
        # the first run when --flatten let a mixed-font paragraph through.
        text_runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
        main = max(text_runs, key=lambda r: sum(len(t.text or "") for t in r.findall(qn("w:t"))), default=None)
        base = main.find(qn("w:rPr")) if main is not None else None
        if flatten and len({base_key(r.find(qn("w:rPr"))) for r in text_runs}) > 1:
            warnings.append(f"[{i}] flattened: mixed fonts or sizes now all match the main text")
        base = copy.deepcopy(base) if base is not None else None
        for child in list(p):
            if child.tag in (qn("w:r"), qn("w:proofErr")):
                p.remove(child)
        for text, bold, italic in segs:
            p.append(make_run(base, text, bold, italic))
        done += 1
    doc.save(dst)
    print(f"rewrote {done} of {len(edits)} paragraphs -> {dst}")
    total = sum(len(run_text(p, marks=False).split()) for p in paragraphs(Document(src)))
    print(f"words in edited paragraphs: {words_old} -> {words_new}; "
          f"whole document: {total} -> {total - words_old + words_new}")
    for w in warnings:
        print(w)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) == 4 and sys.argv[1] == "extract":
        extract(sys.argv[2], sys.argv[3])
    elif len(sys.argv) in (5, 6) and sys.argv[1] == "apply" and sys.argv[5:] in ([], ["--flatten"]):
        apply(sys.argv[2], sys.argv[3], sys.argv[4], flatten=len(sys.argv) == 6)
    else:
        sys.exit(__doc__)
